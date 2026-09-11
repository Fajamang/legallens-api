import io
import json
import sys
import sqlite3
from pathlib import Path
import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as backend

TEXT = 'Een overeenkomst tussen huurder en verhuurder met afspraken over betaling en opzegging. '
RESULT = {'summary': 'De overeenkomst bevat afspraken.', 'contract_type': 'Huurovereenkomst',
          'parties_involved': ['Huurder', 'Verhuurder'], 'key_dates': {}, 'risks': [],
          'overall_advice': 'Controleer de opzegtermijn.', 'sentiment_score': .5,
          'action_plan': {'direct': ['Controleer de termijn']}, 'negotiation_strategy': {},
          'due_diligence_findings': [], 'time_saved_hours': 0}

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(backend, 'DB_PATH', tmp_path / 'data.db')
    monkeypatch.setattr(backend, 'UPLOAD_DIR', tmp_path / 'uploads')
    monkeypatch.setattr(backend, 'VALID_API_KEYS', ['local-integration-key'])
    with TestClient(backend.app, headers={'X-API-Key': 'local-integration-key'}) as c:
        yield c

@pytest.fixture
def ai(monkeypatch):
    async def answer(*args, **kwargs):
        return dict(RESULT)
    monkeypatch.setattr(backend.analyzer, 'analyze_text', answer)


def test_auth_and_static(client):
    assert client.get('/').status_code == 200
    assert client.get('/static/app.js').status_code == 200
    assert client.get('/health').json()['status'] == 'healthy'
    assert client.get('/api/dossiers', headers={'X-API-Key': ''}).status_code == 401
    assert client.get('/api/dossiers', headers={'X-API-Key': 'wrong'}).status_code == 401
    assert client.get('/api/status').status_code == 200


def test_fail_closed(client, monkeypatch):
    monkeypatch.setattr(backend, 'VALID_API_KEYS', [])
    assert client.get('/api/dossiers').status_code == 503


def test_dossier_file_crud_and_cascade(client, ai):
    d = client.post('/api/dossiers', json={'name': 'Zaak', 'client': 'Klant'}).json()
    id = d['id']
    assert client.put('/api/dossiers/'+id, json={'name': 'Zaak bijgewerkt', 'status': 'archived'}).status_code == 200
    f = client.post(f'/api/dossiers/{id}/files', files={'file': ('../TEST.TXT', TEXT.encode())}).json()
    assert f['original_name'] == 'TEST.TXT'
    assert client.get('/api/files/'+f['id']).content == TEXT.encode()
    a = client.post('/api/analyze-text', json={'text': TEXT, 'dossier_id': id, 'mode': 'advocaat'}).json()
    assert client.get('/api/dossiers/'+id).json()['analyses'][0]['id'] == a['document_id']
    assert client.get('/api/stats').json()['total_files'] == 1
    assert client.delete('/api/dossiers/'+id).status_code == 200
    assert client.get('/api/files/'+f['id']).status_code == 404
    assert client.get('/api/analyses/'+a['document_id']).json()['dossier_id'] is None
    assert not list(backend.UPLOAD_DIR.iterdir())
    assert client.delete('/api/dossiers/'+id).status_code == 404


def test_analysis_validation_persistence_and_search(client, ai):
    assert client.post('/api/analyze-text', json={'text': TEXT, 'dossier_id': 'missing'}).status_code == 404
    assert client.post('/api/analyze-text', json={'text': 'short'}).status_code == 422
    assert client.post('/api/analyze-text', json={'text': TEXT, 'mode': 'invalid'}).status_code == 422
    a = client.post('/api/analyze-text', json={'text': TEXT, 'analysis_type': 'due_diligence'}).json()
    assert client.get('/api/analyses?q=overeenkomst').json()['analyses'][0]['document_id'] == a['document_id']
    assert not client.get('/api/analyses?q=no-match').json()['analyses']
    assert client.delete('/api/analyses/'+a['document_id']).status_code == 200
    assert client.get('/api/analyses/'+a['document_id']).status_code == 404


def test_bad_ai_not_saved(client, monkeypatch):
    async def invalid(*args): return {'summary': 'Incomplete'}
    monkeypatch.setattr(backend.analyzer, 'analyze_text', invalid)
    assert client.post('/api/analyze-text', json={'text': TEXT}).status_code == 502
    assert not client.get('/api/analyses').json()['analyses']


def test_missing_ai_does_not_mock(client, monkeypatch):
    monkeypatch.setattr(backend, 'OPENAI_API_KEY', '')
    assert client.post('/api/analyze-text', json={'text': TEXT}).status_code == 503
    assert not client.get('/api/analyses').json()['analyses']


def test_txt_docx_and_pdf(client, ai):
    doc = Document(); doc.add_paragraph(TEXT)
    table = doc.add_table(rows=1, cols=1); table.cell(0, 0).text = 'Bepaling uit een tabel.'
    stream = io.BytesIO(); doc.save(stream)
    from reportlab.pdfgen import canvas
    pdf = io.BytesIO(); c = canvas.Canvas(pdf); c.drawString(40, 750, TEXT); c.save()
    for name, content in [('TEST.TXT', TEXT.encode()), ('test.docx', stream.getvalue()), ('test.pdf', pdf.getvalue())]:
        r = client.post('/api/analyze-file', files={'file': (name, content)})
        assert r.status_code == 200, r.text
        assert client.get('/api/analyses/'+r.json()['document_id']).status_code == 200
    assert len(client.get('/api/analyses').json()['analyses']) == 3
    assert 'Bepaling uit een tabel.' in backend.extract_text(stream.getvalue(), '.docx')


def test_upload_errors(client, ai, monkeypatch):
    assert client.post('/api/analyze-file', files={'file': ('x.exe', b'x')}).status_code == 400
    assert client.post('/api/analyze-file', files={'file': ('x.pdf', b'broken')}).status_code == 400
    assert client.post('/api/analyze-file', files={'file': ('x.docx', b'broken')}).status_code == 400
    assert client.post('/api/analyze-file', files={'file': ('x.txt', b'short')}).status_code == 400
    assert client.post('/api/analyze-file', files={'file': ('x.txt', b'')}).status_code == 400
    monkeypatch.setattr(backend, 'MAX_UPLOAD_BYTES', 20)
    assert client.post('/api/analyze-file', files={'file': ('x.txt', TEXT.encode())}).status_code == 413


def test_template_generate_exports_delete(client):
    templates = client.get('/api/templates').json()['templates']
    for template in templates:
        assert client.post('/api/templates/generate', json={'template_id': template['id']}).status_code == 422
        d = client.post('/api/templates/generate', json={'template_id': template['id'], 'custom_fields': {f: 'Ingevuld & <tekst>' for f in template['fields']}}).json()
        for format in ['txt', 'docx', 'pdf']:
            r = client.get(f"/api/documents/{d['document_id']}/download?format={format}")
            assert r.status_code == 200
            if format == 'docx': assert 'Ingevuld' in '\n'.join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
            if format == 'pdf': assert r.content.startswith(b'%PDF')
        assert client.delete('/api/documents/'+d['document_id']).status_code == 200
    assert client.post('/api/templates/generate', json={'template_id':'unknown'}).status_code == 404


def test_search_and_commentary_sources(client, monkeypatch):
    monkeypatch.setattr(backend, 'TAVILY_API_KEY', 'local-search-key')
    async def remote(url, payload, headers, timeout):
        assert headers['Authorization'] == 'Bearer local-search-key'
        domain = payload['include_domains'][0]
        return {'results': [{'title':'Artikel 6:74 BW','url':'https://'+domain+'/example','content':'Bronfragment'},
                            {'url':'https://evil.example','content':'Uitsluiten'}]}
    async def commentary(*args): return {'commentary': 'Op basis van de fragmenten.'}
    monkeypatch.setattr(backend, 'remote_json', remote)
    monkeypatch.setattr(backend, 'openai_json', commentary)
    r = client.post('/api/legal-commentary', json={'article':'6:74 BW'})
    assert r.status_code == 200
    assert len(r.json()['sources']) == 1
    assert r.json()['text_kind'] == 'search_snippets'
    assert len(r.json()['jurisprudence']) == 1


def test_openai_payload_and_errors(client, monkeypatch):
    monkeypatch.setattr(backend, 'OPENAI_API_KEY', 'local-openai-test')
    async def remote(url, payload, headers, timeout=90):
        request = json.loads(payload['messages'][1]['content'])
        assert request['mode'] == 'advocaat'
        assert request['document'] == TEXT * 70  # exceeds old silent 4,000-character cutoff
        return {'choices':[{'finish_reason':'stop','message':{'content':json.dumps(RESULT)}}]}
    monkeypatch.setattr(backend, 'remote_json', remote)
    assert client.post('/api/analyze-text', json={'text':TEXT*70, 'mode':'advocaat'}).status_code == 200
    async def incomplete(*args, **kwargs): return {'choices':[{'finish_reason':'length','message':{'content':'{}'}}]}
    monkeypatch.setattr(backend, 'remote_json', incomplete)
    assert client.post('/api/analyze-text', json={'text':TEXT}).status_code == 502
    assert len(client.get('/api/analyses').json()['analyses']) == 1


def test_legacy_database_is_preserved(tmp_path, client):
    # init_db is idempotent and keeps existing rows in the original v9 schema.
    id = client.post('/api/dossiers', json={'name':'Bestaand dossier'}).json()['id']
    backend.init_db()
    assert client.get('/api/dossiers/'+id).json()['name'] == 'Bestaand dossier'


def test_request_cap_and_security_headers(client, monkeypatch):
    monkeypatch.setattr(backend, 'MAX_UPLOAD_BYTES', 20)
    r = client.post('/api/analyze-file', content=b'x'*(1024*1024+21), headers={'Content-Type':'application/octet-stream'})
    assert r.status_code == 413
    assert client.get('/api/stats').headers['cache-control'] == 'no-store'
    assert "script-src 'self'" in client.get('/').headers['content-security-policy']


CASE_XML = b'''<open-rechtspraak xmlns:dc="http://purl.org/dc/terms/">
<dc:identifier>ECLI:NL:HR:2021:1778</dc:identifier><dc:creator>Hoge Raad</dc:creator>
<dc:date>2021-11-26</dc:date><dc:title>Testuitspraak</dc:title>
<inhoudsindicatie><para>Testinhoudsindicatie.</para></inhoudsindicatie>
<uitspraak><section><title>Beoordeling</title><para>Gepubliceerde tekst.</para></section></uitspraak>
</open-rechtspraak>'''


def test_official_xml_and_metadata_only():
    result = backend.parse_rechtspraak(CASE_XML, 'ECLI:NL:HR:2021:1778')
    assert result['court'] == 'Hoge Raad'
    assert result['date'] == '2021-11-26'
    assert result['full_text'] == 'Beoordeling\n\nGepubliceerde tekst.'
    assert result['metadata_verified']
    metadata = CASE_XML[:CASE_XML.index(b'<uitspraak>')] + b'</open-rechtspraak>'
    assert not backend.parse_rechtspraak(metadata, 'ECLI:NL:HR:2021:1778')['full_text_available']


def test_unsafe_or_mismatched_xml():
    from fastapi import HTTPException
    with pytest.raises(HTTPException):
        backend.parse_rechtspraak(b'<!DOCTYPE doc [<!ENTITY x "expanded">]><doc>&x;</doc>', 'ECLI:NL:HR:2021:1778')
    with pytest.raises(HTTPException):
        backend.parse_rechtspraak(CASE_XML, 'ECLI:NL:HR:2000:1234')


def test_direct_ecli_without_paid_services(client, monkeypatch):
    monkeypatch.setattr(backend, 'TAVILY_API_KEY', '')
    monkeypatch.setattr(backend, 'OPENAI_API_KEY', '')
    async def case(ecli):
        return backend.parse_rechtspraak(CASE_XML, backend.clean_ecli(ecli))
    monkeypatch.setattr(backend.rechtspraak, 'get_case', case)
    result = client.post('/api/legal-commentary', json={'article':'ecli:nl:hr:2021:1778'})
    assert result.status_code == 200
    assert result.json()['text_kind'] == 'official_case'
    assert result.json()['jurisprudence'][0]['full_text_available']
    assert client.get('/api/rechtspraak/ECLI:NL:HR:2021:1778').status_code == 200
    assert client.get('/api/rechtspraak/invalid').status_code == 422
    assert client.post('/api/legal-commentary', json={'article':'6:74 BW'}).status_code == 503


def test_tavily_enrichment_and_deduplication(client, monkeypatch):
    monkeypatch.setattr(backend, 'TAVILY_API_KEY', 'local-search')
    lookups = []
    async def search(url, payload, headers, timeout):
        if 'wetten.overheid.nl' in payload['include_domains']: return {'results':[]}
        return {'results':[{'title':'Testuitspraak', 'url':'https://uitspraken.rechtspraak.nl/details?id=ECLI:NL:HR:2021:1778','content':'Snippet'}]*2}
    async def case(ecli):
        lookups.append(ecli)
        return backend.parse_rechtspraak(CASE_XML, ecli)
    async def ai(*args): return {'commentary':'Commentaar op fragmenten.'}
    monkeypatch.setattr(backend, 'remote_json', search)
    monkeypatch.setattr(backend, 'openai_json', ai)
    monkeypatch.setattr(backend.rechtspraak, 'get_case', case)
    r = client.post('/api/legal-commentary', json={'article':'6:74 BW'}).json()
    assert len(r['jurisprudence']) == 1
    assert lookups == ['ECLI:NL:HR:2021:1778']
    assert r['jurisprudence'][0]['provider'] == 'rechtspraak_open_data'
    assert r['jurisprudence'][0]['found_via'] == 'tavily'


def test_official_api_failure_keeps_tavily_result(client, monkeypatch):
    from fastapi import HTTPException
    monkeypatch.setattr(backend, 'TAVILY_API_KEY', 'local-search')
    monkeypatch.setattr(backend, 'OPENAI_API_KEY', '')
    async def search(url, payload, headers, timeout):
        if 'wetten.overheid.nl' in payload['include_domains']: return {'results':[]}
        return {'results':[{'title':'ECLI:NL:HR:2021:1778','url':'https://uitspraken.rechtspraak.nl/details?id=ECLI:NL:HR:2021:1778','content':'Fallbackfragment'}]}
    async def failed(ecli): raise HTTPException(504, 'Teststoring Rechtspraak')
    monkeypatch.setattr(backend, 'remote_json', search)
    monkeypatch.setattr(backend.rechtspraak, 'get_case', failed)
    r = client.post('/api/legal-commentary', json={'article':'6:74 BW'})
    assert r.status_code == 200
    assert r.json()['jurisprudence'][0]['provider'] == 'tavily'
    assert not r.json()['jurisprudence'][0]['metadata_verified']
    assert 'Fallbackfragment' == r.json()['jurisprudence'][0]['content']
    assert len(r.json()['warnings']) == 2  # Rechtspraak and unavailable OpenAI


def test_official_request_cache_and_url(monkeypatch):
    import asyncio
    original_client = httpx.AsyncClient
    calls = []
    def handler(request):
        calls.append(request)
        assert request.url.host == 'data.rechtspraak.nl'
        assert request.url.path == '/uitspraken/content'
        assert request.url.params['id'] == 'ECLI:NL:HR:2021:1778'
        assert 'authorization' not in request.headers
        return httpx.Response(200, content=CASE_XML)
    monkeypatch.setattr(backend.httpx, 'AsyncClient', lambda **kw: original_client(transport=httpx.MockTransport(handler), **kw))
    async def run():
        c = backend.RechtspraakClient()
        a = await c.get_case('ECLI:NL:HR:2021:1778')
        b = await c.get_case('ECLI:NL:HR:2021:1778')
        assert a == b
    asyncio.run(run())
    assert len(calls) == 1
